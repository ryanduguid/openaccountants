"""Output and coverage regressions for the local tooling optimisations."""

import importlib.util
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))


def load(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class LocalOptimisationTests(unittest.TestCase):
    def test_validation_ranges_cover_only_original_body_cells_after_round_trip(self):
        import openpyxl

        generator = load("range_generator", "tools/workbooks/generate_workbook.py")
        workbook = openpyxl.Workbook()
        title = generator._add_skill_tab(workbook, {"tab": "Synthetic", "sections": [
            {"section": "Empty", "rows": []},
            {"section": "First", "rows": [{"item": "A"}, {"item": "B"}]},
            {"section": "Second", "rows": [{"item": "C"}]},
        ]})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "synthetic.xlsx"
            workbook.save(path)
            restored = openpyxl.load_workbook(path)
            sheet = restored[title]
            validation = next(iter(sheet.data_validations.dataValidation))
            covered = {row for row in range(1, sheet.max_row + 1) if f"C{row}" in validation}
            self.assertEqual(covered, {4, 5, 7, *range(11, 19)})
            self.assertEqual(len(validation.sqref.ranges), 3)
            restored.close()

    def test_multiple_concepts_reuse_extraction_without_changing_each_claim(self):
        scanner = load("cached_scanner", "scripts/detect-contradictions.py")
        spec = {"terms": [re.compile("rate")], "require": None, "exclude": None,
                "kind": "any"}
        compiled = {"US": {"first": spec, "second": spec}}
        text = "---\nname: synthetic\njurisdiction: US\ncategory: federal\ntax_year: 2025\n---\nrate 10% and $100 in 2025\n"
        stats = {"claims": 0, "multivalue_lines_skipped": 0,
                 "ambiguous_year_dropped": 0, "historical_dropped": 0}
        with mock.patch.object(scanner, "extract_percent_values", wraps=scanner.extract_percent_values) as percent, \
                mock.patch.object(scanner, "extract_money_values", wraps=scanner.extract_money_values) as money, \
                mock.patch.object(scanner, "sentence_years", wraps=scanner.sentence_years) as years:
            claims = scanner.extract_claims("synthetic.md", text, "US", compiled, stats, 2025)
        self.assertEqual((percent.call_count, money.call_count, years.call_count), (1, 1, 1))
        self.assertEqual([(claim["concept"], claim["kind"], claim["value"], claim["year"]) for claim in claims],
                         [("first", "money", 100.0, 2025), ("first", "percent", 10.0, 2025),
                          ("second", "money", 100.0, 2025), ("second", "percent", 10.0, 2025)])

    def test_index_freshness_reports_generation_failure_and_ignores_generation_time(self):
        validator = load("direct_index_validator", "scripts/validate-guides.py")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            builder = root / "builder.py"
            index = {"counts": {"guides": 1}, "guides": [{"slug": "synthetic"}]}
            (root / "index.json").write_text(json.dumps({**index, "generated_at": "old"}))
            builder.write_text("def build_index():\n    return " + repr({**index, "generated_at": "new"}))
            with mock.patch.object(validator, "REPO_ROOT", folder), mock.patch.object(validator, "BUILD_INDEX", str(builder)), \
                    mock.patch.object(validator.subprocess, "run", side_effect=AssertionError("Unexpected child process")):
                errors = []
                validator.check_index_fresh(errors)
                self.assertEqual(errors, [])
                builder.write_text("raise RuntimeError('synthetic failure')\n")
                validator.check_index_fresh(errors)
                self.assertIn("synthetic failure", errors[0])


if __name__ == "__main__":
    unittest.main()
